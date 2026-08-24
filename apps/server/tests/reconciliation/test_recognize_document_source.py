"""Naming a document by hand when the classifier could not.

The exogena report is read by a parser rather than by OCR against a configured
document type, so it always fails classification. These cover the path that
lets a person say what it is and get it to a state that can be approved.
"""

from __future__ import annotations

import io
from datetime import UTC, datetime

import fixtures
import pytest
from openpyxl import Workbook

from server.application.use_cases import (
    ApproveDocument,
    ApproveDocumentInput,
    DocumentAlreadyApproved,
    DocumentNotFound,
    DocumentNotRecognized,
    RecognizeDocumentSource,
    RecognizeDocumentSourceInput,
)
from server.domain.entities import Document, DocumentStatus, ExtractedData
from server.domain.ports import DocumentContent
from server.infrastructure.adapters.in_memory_repositories import (
    InMemoryDocumentRepository,
    InMemoryExtractedDataRepository,
)
from server.reconciliation.core.registry import KindRegistry
from server.reconciliation.infrastructure import KindSourceParsers
from server.reconciliation.kinds.exogena import ExogenaReconciliation

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXOGENA_SOURCE = "exogena_report"


class _Storage:
    """DocumentStorage stub: hands back one file whatever is asked for."""

    def __init__(self, data: bytes, mime_type: str = XLSX) -> None:
        self._content = DocumentContent(
            data=data, mime_type=mime_type, file_name="reporteExogena2025.xlsx"
        )

    def download(self, file_reference: str) -> DocumentContent:
        return self._content

    def list_files(self, folder_reference: str):
        return []


def _failed_document(mime_type: str = XLSX) -> Document:
    """A document exactly as intake leaves the exogena: FAILED, no type."""
    return Document(
        id="doc-1",
        client_id="client-1",
        document_type_id=None,
        drive_file_id="drive-1",
        file_name="reporteExogena2025.xlsx",
        mime_type=mime_type,
        status=DocumentStatus.FAILED,
        error="Could not identify the document type",
        created_at=datetime.now(UTC),
    )


def _unrelated_spreadsheet() -> bytes:
    """An ordinary workbook. A client's folder holds all sorts of them."""
    workbook = Workbook()
    workbook.active.append(["Fecha", "Concepto", "Monto"])
    workbook.active.append(["2025-01-02", "Arriendo", 1_200_000])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _use_case(
    documents: InMemoryDocumentRepository,
    extracted: InMemoryExtractedDataRepository,
    storage: _Storage,
) -> RecognizeDocumentSource:
    return RecognizeDocumentSource(
        documents=documents,
        storage=storage,
        parsers=KindSourceParsers(KindRegistry([ExogenaReconciliation()])),
        extracted_data=extracted,
    )


def test_a_failed_exogena_becomes_processed_and_carries_its_source() -> None:
    documents = InMemoryDocumentRepository()
    documents.save(_failed_document())
    extracted = InMemoryExtractedDataRepository()
    use_case = _use_case(documents, extracted, _Storage(fixtures.exogena_workbook_bytes()))

    recognized = use_case.execute(
        RecognizeDocumentSourceInput(document_id="doc-1", source_id=EXOGENA_SOURCE)
    )

    assert recognized.document.status == DocumentStatus.PROCESSED
    assert recognized.document.source_id == EXOGENA_SOURCE
    assert recognized.document.error is None
    assert recognized.document.processed_at is not None
    # No document type, deliberately: this file is not extracted against one.
    assert recognized.document.document_type_id is None
    assert documents.get("doc-1") == recognized.document
    # The file states the year it covers, and it travels with the document so a
    # caller knows which reports just went stale.
    assert recognized.periods == ("2025",)


def test_what_was_read_is_stored_as_a_summary_not_as_rows() -> None:
    """The report runs to thousands of rows and they already reach
    reconciliation from the file. What a reviewer needs is enough to tell the
    right file was read."""
    documents = InMemoryDocumentRepository()
    documents.save(_failed_document())
    extracted = InMemoryExtractedDataRepository()
    use_case = _use_case(documents, extracted, _Storage(fixtures.exogena_workbook_bytes()))

    use_case.execute(RecognizeDocumentSourceInput(document_id="doc-1", source_id=EXOGENA_SOURCE))

    stored = extracted.get_by_document("doc-1")
    assert stored is not None
    assert stored.fields["reported_rows"] == len(fixtures.EXOGENA_ROWS)
    assert stored.fields["periods"] == ["2025"]
    assert stored.fields["total_reported"] == sum(row[3] for row in fixtures.EXOGENA_ROWS)
    # A parser reads the file exactly or fails: reporting a confidence would
    # put it on the same scale as an AI's guess.
    assert stored.confidence is None


def test_a_recognized_document_can_then_be_approved_and_keeps_its_source() -> None:
    """The whole point: the exogena reaches a state a person can sign off on,
    and the approval does not erase what it was read as."""
    documents = InMemoryDocumentRepository()
    documents.save(_failed_document())
    extracted = InMemoryExtractedDataRepository()
    _use_case(documents, extracted, _Storage(fixtures.exogena_workbook_bytes())).execute(
        RecognizeDocumentSourceInput(document_id="doc-1", source_id=EXOGENA_SOURCE)
    )

    approved = ApproveDocument(documents).execute(
        ApproveDocumentInput(document_id="doc-1", approved_by="jane")
    )

    assert approved.status == DocumentStatus.APPROVED
    assert approved.source_id == EXOGENA_SOURCE


def test_recognizing_replaces_an_earlier_extraction_rather_than_adding_one() -> None:
    documents = InMemoryDocumentRepository()
    documents.save(_failed_document())
    extracted = InMemoryExtractedDataRepository()
    extracted.save(
        ExtractedData(
            id="extraction-1",
            document_id="doc-1",
            fields={"stale": True},
            confidence=0.4,
            created_at=datetime.now(UTC),
        )
    )
    use_case = _use_case(documents, extracted, _Storage(fixtures.exogena_workbook_bytes()))

    use_case.execute(RecognizeDocumentSourceInput(document_id="doc-1", source_id=EXOGENA_SOURCE))

    stored = extracted.get_by_document("doc-1")
    assert stored is not None
    assert stored.id == "extraction-1"
    assert "stale" not in stored.fields


def test_a_file_that_is_not_the_named_report_leaves_the_document_untouched() -> None:
    """Choosing the wrong source must leave no trace, or the next attempt would
    start from a document the first one damaged."""
    documents = InMemoryDocumentRepository()
    documents.save(_failed_document())
    extracted = InMemoryExtractedDataRepository()
    use_case = _use_case(documents, extracted, _Storage(_unrelated_spreadsheet()))

    with pytest.raises(DocumentNotRecognized):
        use_case.execute(
            RecognizeDocumentSourceInput(document_id="doc-1", source_id=EXOGENA_SOURCE)
        )

    untouched = documents.get("doc-1")
    assert untouched.status == DocumentStatus.FAILED
    assert untouched.source_id is None
    assert untouched.processed_at is None
    assert extracted.get_by_document("doc-1") is None


def test_a_pdf_is_refused_before_it_is_ever_parsed() -> None:
    documents = InMemoryDocumentRepository()
    documents.save(_failed_document(mime_type="application/pdf"))
    extracted = InMemoryExtractedDataRepository()
    use_case = _use_case(documents, extracted, _Storage(b"%PDF-1.4", mime_type="application/pdf"))

    with pytest.raises(DocumentNotRecognized):
        use_case.execute(
            RecognizeDocumentSourceInput(document_id="doc-1", source_id=EXOGENA_SOURCE)
        )


def test_an_unknown_source_is_refused() -> None:
    documents = InMemoryDocumentRepository()
    documents.save(_failed_document())
    use_case = _use_case(
        documents,
        InMemoryExtractedDataRepository(),
        _Storage(fixtures.exogena_workbook_bytes()),
    )

    with pytest.raises(DocumentNotRecognized):
        use_case.execute(RecognizeDocumentSourceInput(document_id="doc-1", source_id="nope"))


def test_a_missing_document_is_reported_as_such() -> None:
    use_case = _use_case(
        InMemoryDocumentRepository(),
        InMemoryExtractedDataRepository(),
        _Storage(fixtures.exogena_workbook_bytes()),
    )

    with pytest.raises(DocumentNotFound):
        use_case.execute(
            RecognizeDocumentSourceInput(document_id="missing", source_id=EXOGENA_SOURCE)
        )


def test_an_approved_document_is_never_re_read_behind_the_approval() -> None:
    """Approval is recorded against what was on screen. Replacing the
    extraction underneath would leave it standing over content nobody saw."""
    documents = InMemoryDocumentRepository()
    approved = _failed_document()
    documents.save(
        Document(
            id=approved.id,
            client_id=approved.client_id,
            document_type_id=None,
            drive_file_id=approved.drive_file_id,
            file_name=approved.file_name,
            mime_type=approved.mime_type,
            status=DocumentStatus.APPROVED,
            error=None,
            created_at=approved.created_at,
            approved_by="jane",
            source_id=EXOGENA_SOURCE,
        )
    )
    use_case = _use_case(
        documents, InMemoryExtractedDataRepository(), _Storage(fixtures.exogena_workbook_bytes())
    )

    with pytest.raises(DocumentAlreadyApproved):
        use_case.execute(
            RecognizeDocumentSourceInput(document_id="doc-1", source_id=EXOGENA_SOURCE)
        )


def test_a_recognized_exogena_can_be_approved_over_http() -> None:
    """End to end, the flow a reviewer actually walks: open the document, say
    what it is, approve it. Lives beside the exogena fixture rather than with
    the other API tests, because rebuilding the DIAN's layout there would mean
    maintaining a second copy of it."""
    from fastapi.testclient import TestClient

    from server.infrastructure.api import deps
    from server.infrastructure.api.auth_dependency import require_session
    from server.infrastructure.api.deps import get_document_repository
    from server.main import app

    get_document_repository.cache_clear()
    documents = get_document_repository()
    documents.save(_failed_document())
    storage = _Storage(fixtures.exogena_workbook_bytes())
    original_storage = deps.get_document_storage
    deps.get_document_storage = lambda: storage
    app.dependency_overrides[require_session] = lambda: None
    try:
        client = TestClient(app)

        recognized = client.post("/documents/doc-1/recognize", json={"source_id": EXOGENA_SOURCE})
        assert recognized.status_code == 200
        assert recognized.json()["status"] == "processed"
        assert recognized.json()["source_id"] == EXOGENA_SOURCE

        approved = client.post("/documents/doc-1/approve", json={"approved_by": "jane"})
        assert approved.status_code == 200
        assert approved.json()["status"] == "approved"
        assert approved.json()["source_id"] == EXOGENA_SOURCE
    finally:
        deps.get_document_storage = original_storage
        app.dependency_overrides.clear()
        get_document_repository.cache_clear()
